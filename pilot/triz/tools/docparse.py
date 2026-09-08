"""Local extraction with page/sheet/row provenance and bounded OCR."""
import csv
import io
import os
import re
from pathlib import Path
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
ALLOWED_EXT = TEXT_EXT | IMAGE_EXT | {".pdf", ".xlsx"}
def kind_of(filename):
    ext = Path(filename).suffix.lower()
    return "REPORT" if ext == ".pdf" else "DRAWING" if ext in IMAGE_EXT else "DATA" if ext in {".csv", ".xlsx"} else "SPEC"
def ocr(image):
    import pytesseract
    try:
        return pytesseract.image_to_string(image, lang=os.getenv("OCR_LANG", "eng"), timeout=20)
    except (pytesseract.TesseractNotFoundError, RuntimeError) as exc:
        return "[미추출: OCR 실행 파일 또는 언어팩 확인 필요. 도면 관계는 텍스트로 보완해 주세요.]"
def extract(path: Path, filename=""):
    name = filename or path.name
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ValueError("지원하지 않는 첨부 형식입니다.")
    lines = []
    def add(location, text):
        for line in str(text or "").splitlines():
            if line.strip():
                lines.append(f"[{name} · {location}] {line.strip()[:1200]}")
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages[:30], 1):
                text = page.extract_text() or ""
                if not text.strip():
                    try:
                        text = ocr(page.to_image(resolution=130).original)
                    except Exception:
                        text = "[미추출: 스캔 페이지. 원본 확인과 설명이 필요합니다.]"
                add(f"p.{i}", text)
                for t, rows in enumerate(page.extract_tables()[:4], 1):
                    for r, row in enumerate(rows[:40], 1):
                        add(f"p.{i} 표{t} 행{r}", " | ".join(str(v or "") for v in row))
            if len(pdf.pages) > 30:
                add("범위", "[일부 추출: 첫 30페이지만 분석]")
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets[:8]:
                for i, row in enumerate(sheet.iter_rows(max_row=200, max_col=25, values_only=True), 1):
                    add(f"{sheet.title} 행{i}", " | ".join(str(v) if v is not None else "" for v in row))
        finally:
            book.close()
        add("범위", "[일부 추출: 최대 8시트·200행·25열. 수식은 저장된 계산값을 사용]")
    elif ext in IMAGE_EXT:
        from PIL import Image
        with Image.open(path) as image:
            if image.width * image.height > 25_000_000:
                raise ValueError("이미지 해상도가 너무 큽니다.")
            add("OCR", ocr(image))
            add("주의", "OCR은 문자 추출입니다. 도면의 연결·치수·기호 관계는 사용자 확인이 필요합니다.")
    else:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp949", errors="replace")
        if ext == ".csv":
            for i, row in enumerate(csv.reader(io.StringIO(text)), 1):
                if i > 200:
                    add("범위", "[일부 추출: 첫 200행]")
                    break
                add(f"행{i}", " | ".join(row[:25]))
        else:
            for i, line in enumerate(text.splitlines()[:1200], 1):
                add(f"행{i}", line)
    # Select measurements and problem-bearing lines across the whole document, not just page 1.
    priority = re.compile(r"\d\s*(nm|mm|um|μm|℃|°C|Pa|MPa|rpm|%|ms|V|A)|불량|제약|결함|온도|압력|failure|defect", re.I)
    ranked = sorted(enumerate(lines), key=lambda pair: (not bool(priority.search(pair[1])), pair[0]))
    facts = [line for _, line in sorted(ranked[:60])]
    selected = "\n".join(lines)
    if len(selected) > 40000:
        selected = selected[:40000] + "\n[일부 추출: 텍스트 길이 제한]"
    return selected, facts
