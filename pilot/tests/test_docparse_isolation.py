"""Bound parser work without large files, paid calls or production storage."""
import json
import os
from pathlib import Path
import subprocess
import sys
import threading
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace

import pytest

from triz.tools import docparse, docparse_worker


@pytest.fixture(autouse=True)
def parser_slots(monkeypatch):
    monkeypatch.setattr(docparse, "_WORKER_SLOTS", None)
    monkeypatch.delenv("DOCUMENT_PARSE_CONCURRENCY", raising=False)


@pytest.mark.parametrize("kind", ["txt", "csv", "xlsx", "pdf", "png"])
def test_normal_extract_matches_local_contract(tmp_path, kind):
    path = tmp_path / ("sample." + kind)
    if kind == "txt":
        path.write_text("측정 결과\n압력 20 MPa\n", encoding="utf-8")
    elif kind == "csv":
        path.write_text("pressure,time\n20 MPa,30 ms\n", encoding="utf-8")
    elif kind == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "process"
        book.active.append(["pressure", "20 MPa"])
        book.save(path)
        book.close()
    elif kind == "pdf":
        from reportlab.pdfgen.canvas import Canvas
        document = Canvas(str(path))
        document.drawString(50, 700, "Pressure 20 MPa")
        document.save()
    else:
        from PIL import Image
        Image.new("RGB", (40, 40), "white").save(path)
    before = path.read_bytes()
    expected = docparse._extract_local(path, "측정." + kind)
    assert docparse.extract(path, "측정." + kind) == expected
    assert path.read_bytes() == before
    assert len(expected[0]) <= 40100 and len(expected[1]) <= 60
    if kind != "png":
        assert "20 MPa" in expected[0]
    else:
        assert "OCR은 문자 추출입니다" in expected[0]


@pytest.mark.parametrize("platform", ["posix", "nt"])
def test_timeout_terminates_worker_and_preserves_source(tmp_path, monkeypatch, platform):
    path = tmp_path / "safe.txt"
    path.write_text("keep", encoding="utf-8")
    calls = []
    class Worker:
        pid = 12345
        returncode = None
        def communicate(self, input=None, timeout=None):
            calls.append((input, timeout))
            if timeout is not None:
                raise subprocess.TimeoutExpired("fixed-worker", timeout)
            self.returncode = -9
        def kill(self):
            calls.append("killed")
    def spawn(command, **kwargs):
        assert command == [sys.executable, "-I", str(Path(docparse.__file__).with_name("docparse_worker.py"))]
        assert "shell" not in kwargs
        assert kwargs["stderr"] == subprocess.DEVNULL
        assert kwargs["start_new_session"] == (platform == "posix")
        return Worker()
    monkeypatch.setattr(docparse.subprocess, "Popen", spawn)
    monkeypatch.setenv("DOCUMENT_PARSE_TIMEOUT_SECONDS", "1")
    monkeypatch.setenv("DOCUMENT_PARSE_MEMORY_MB", "64")
    monkeypatch.setattr(docparse, "os", SimpleNamespace(name=platform, environ=os.environ, getenv=os.getenv,
        killpg=lambda pid, sig: calls.append(("group_killed", pid))))
    monkeypatch.setattr(docparse, "signal", SimpleNamespace(SIGKILL=9))
    def kill_tree(command, **kwargs):
        assert Path(command[0]).name == "taskkill.exe"
        assert command[1:] == ["/PID", "12345", "/T", "/F"]
        assert kwargs["shell"] is False
        assert kwargs["creationflags"] == getattr(subprocess, "CREATE_NO_WINDOW", 0)
        assert kwargs["stdout"] == kwargs["stderr"] == subprocess.DEVNULL
        assert kwargs["timeout"] == 10
        calls.append(("tree_killed", 12345))
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(docparse.subprocess, "run", kill_tree)
    with pytest.raises(ValueError, match="시간이 제한"):
        docparse.extract(path)
    assert calls[0][1] == 1
    assert json.loads(calls[0][0])["path"] == str(path.resolve())
    assert calls[-1] == (None, None)  # reap after killing
    assert (("group_killed", 12345) if platform == "posix" else ("tree_killed", 12345)) in calls
    assert path.read_text(encoding="utf-8") == "keep"


@pytest.mark.parametrize("name,value", [
    ("DOCUMENT_PARSE_TIMEOUT_SECONDS", "0"),
    ("DOCUMENT_PARSE_TIMEOUT_SECONDS", "3601"),
    ("DOCUMENT_PARSE_TIMEOUT_SECONDS", "not-a-number"),
    ("DOCUMENT_PARSE_MEMORY_MB", "63"),
    ("DOCUMENT_PARSE_MEMORY_MB", "4097"),
    ("DOCUMENT_PARSE_CONCURRENCY", "0"),
    ("DOCUMENT_PARSE_CONCURRENCY", "5"),
    ("DOCUMENT_PARSE_CONCURRENCY", "invalid"),
])
def test_invalid_limit_configuration_never_starts_worker(tmp_path, monkeypatch, name, value):
    monkeypatch.setenv(name, value)
    monkeypatch.setattr(docparse.subprocess, "Popen", lambda *a, **kw: pytest.fail("worker started"))
    with pytest.raises(ValueError, match="제한 설정"):
        docparse.extract(tmp_path / "safe.txt")


@pytest.mark.parametrize("payload,returncode", [
    (b'{"text":"ok","facts":[]}', 1),
    (b"not json", 0),
    (b"x" * (docparse.MAX_RESULT_BYTES + 1), 0),
    (json.dumps({"text": "ok", "facts": ["x"] * 61}).encode(), 0),
    (json.dumps({"text": "x" * 40101, "facts": []}).encode(), 0),
], ids=["worker_failed", "invalid_json", "too_large", "too_many_facts", "too_long_text"])
def test_bad_worker_result_fails_closed(tmp_path, monkeypatch, payload, returncode):
    def spawn(command, **kwargs):
        kwargs["stdout"].write(payload)
        def communicate(**kw):
            assert kw["timeout"] == 660
        return SimpleNamespace(returncode=returncode, communicate=communicate)
    monkeypatch.delenv("DOCUMENT_PARSE_TIMEOUT_SECONDS", raising=False)
    monkeypatch.setattr(docparse.subprocess, "Popen", spawn)
    with pytest.raises(ValueError):
        docparse.extract(tmp_path / "safe.txt")


@pytest.mark.parametrize("sizes", [
    [1] * (docparse.XLSX_MAX_MEMBERS + 1),
    [docparse.XLSX_MAX_MEMBER_BYTES + 1],
    [docparse.XLSX_MAX_MEMBER_BYTES] * 3,
])
def test_xlsx_limits_check_metadata_before_openpyxl(tmp_path, monkeypatch, sizes):
    class Archive:
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
        def infolist(self):
            return [SimpleNamespace(file_size=size, flag_bits=0) for size in sizes]
    monkeypatch.setattr(docparse, "ZipFile", lambda _: Archive())
    import openpyxl
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **kw: pytest.fail("expanded unsafe ZIP"))
    with pytest.raises(ValueError, match="압축 해제 크기"):
        docparse._extract_local(tmp_path / "safe.xlsx")


@pytest.mark.parametrize("large_box", ["page", "mediabox", "cropbox"])
def test_pdf_pixel_limit_is_checked_before_rendering(tmp_path, monkeypatch, large_box):
    page = SimpleNamespace(width=600, height=800, mediabox=(0, 0, 600, 800), cropbox=(0, 0, 600, 800),
        extract_text=lambda: "", extract_tables=lambda: [],
        to_image=lambda **kw: pytest.fail("rendered oversized page"))
    if large_box == "page":
        page.width = page.height = 10000
    else:
        setattr(page, large_box, (0, 0, 10000, 10000))
    class Document:
        pages = [page]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
    import pdfplumber
    monkeypatch.setattr(pdfplumber, "open", lambda _: Document())
    text, _ = docparse._extract_local(tmp_path / "safe.pdf")
    assert "페이지 해상도가 너무 큽니다" in text


def test_pdf_normal_ocr_uses_same_resolution_and_text(tmp_path, monkeypatch):
    marker = object()
    resolutions = []
    def render(**kwargs):
        resolutions.append(kwargs["resolution"])
        return SimpleNamespace(original=marker)
    page = SimpleNamespace(width=600, height=800, extract_text=lambda: "", extract_tables=lambda: [], to_image=render)
    class Document:
        pages = [page]
        def __enter__(self):
            return self
        def __exit__(self, *args):
            pass
    import pdfplumber
    monkeypatch.setattr(pdfplumber, "open", lambda _: Document())
    def ocr(image):
        assert image is marker
        return "압력 20 MPa"
    monkeypatch.setattr(docparse, "ocr", ocr)
    assert docparse._extract_local(tmp_path / "safe.pdf") == ("[safe.pdf · p.1] 압력 20 MPa", ["[safe.pdf · p.1] 압력 20 MPa"])
    assert resolutions == [130]


@pytest.mark.parametrize("existing", [(-1, -1), (96 * 1024 * 1024, 128 * 1024 * 1024)])
def test_linux_memory_limit_is_applied_without_raising_existing_cap(monkeypatch, existing):
    recorded = []
    fake = SimpleNamespace(RLIMIT_AS=9, RLIM_INFINITY=-1, getrlimit=lambda _: existing,
        setrlimit=lambda kind, limits: recorded.append((kind, limits)))
    monkeypatch.setitem(sys.modules, "resource", fake)
    monkeypatch.setattr(docparse_worker.sys, "platform", "linux")
    monkeypatch.setenv("DOCUMENT_PARSE_MEMORY_MB", "128")
    docparse_worker.apply_limits()
    wanted = 128 * 1024 * 1024 if existing == (-1, -1) else existing[0]
    assert recorded == [(9, (wanted, wanted))]


def test_windows_does_not_import_linux_resource_module(monkeypatch):
    monkeypatch.setattr(docparse_worker.sys, "platform", "win32")
    monkeypatch.setitem(sys.modules, "resource", None)
    docparse_worker.apply_limits()


@pytest.mark.parametrize("configured,capacity", [(None, 2), ("1", 1), ("4", 4)])
def test_simultaneous_requests_share_worker_limit(tmp_path, monkeypatch, configured, capacity):
    if configured is not None:
        monkeypatch.setenv("DOCUMENT_PARSE_CONCURRENCY", configured)
    state = {"active": 0, "peak": 0, "started": 0}
    lock = threading.Lock()
    ready, release = threading.Event(), threading.Event()
    barrier = threading.Barrier(8)
    class Worker:
        returncode = 0
        def communicate(self, **kwargs):
            try:
                assert release.wait(5), "mock worker was not released"
            finally:
                with lock:
                    state["active"] -= 1
    def spawn(command, **kwargs):
        kwargs["stdout"].write(b'{"text":"ok","facts":[]}')
        with lock:
            state["active"] += 1
            state["started"] += 1
            state["peak"] = max(state["peak"], state["active"])
            if state["active"] == capacity:
                ready.set()
        return Worker()
    monkeypatch.setattr(docparse.subprocess, "Popen", spawn)
    def extract():
        barrier.wait(timeout=5)
        return docparse.extract(tmp_path / "safe.txt")
    with ThreadPoolExecutor(max_workers=8) as executor:
        futures = [executor.submit(extract) for _ in range(8)]
        try:
            assert ready.wait(5), "worker capacity was not reached"
            with lock:
                assert state["active"] == capacity
        finally:
            release.set()
        assert [f.result(timeout=5) for f in futures] == [("ok", [])] * 8
    assert state == {"active": 0, "peak": capacity, "started": 8}


@pytest.mark.parametrize("failure", ["spawn", "exit", "timeout", "output"])
def test_worker_failure_releases_its_slot(tmp_path, monkeypatch, failure):
    slots = threading.BoundedSemaphore(1)
    monkeypatch.setattr(docparse, "_WORKER_SLOTS", slots)
    class Worker:
        pid = 12345
        returncode = 1 if failure == "exit" else 0
        def communicate(self, input=None, timeout=None):
            if failure == "timeout" and timeout is not None:
                raise subprocess.TimeoutExpired("fixed-worker", timeout)
    def spawn(command, **kwargs):
        if failure == "spawn":
            raise OSError("mock startup failure")
        kwargs["stdout"].write(b"invalid" if failure == "output" else b'{"text":"ok","facts":[]}')
        return Worker()
    monkeypatch.setattr(docparse.subprocess, "Popen", spawn)
    monkeypatch.setattr(docparse, "_terminate_worker", lambda _: None)
    with pytest.raises((OSError, ValueError)):
        docparse.extract(tmp_path / "safe.txt")
    assert slots.acquire(blocking=False), "failed worker retained its slot"
    slots.release()


def test_slot_wait_timeout_does_not_spawn_or_release_unowned_slot(tmp_path, monkeypatch):
    deadlines = []
    def acquire(*, timeout):
        deadlines.append(timeout)
        return False
    slots = SimpleNamespace(acquire=acquire, release=lambda: pytest.fail("released someone else's slot"))
    monkeypatch.setattr(docparse, "_WORKER_SLOTS", slots)
    monkeypatch.setenv("DOCUMENT_PARSE_TIMEOUT_SECONDS", "1")
    monkeypatch.setattr(docparse.subprocess, "Popen", lambda *a, **kw: pytest.fail("worker started while full"))
    with pytest.raises(ValueError, match="대기 시간이 제한"):
        docparse.extract(tmp_path / "safe.txt")
    assert deadlines == [1]


@pytest.mark.parametrize("failure", ["exit", "timeout", "missing"])
def test_windows_tree_cleanup_failure_still_kills_worker(monkeypatch, failure):
    killed = []
    process = SimpleNamespace(pid=12345, kill=lambda: killed.append(True))
    monkeypatch.setattr(docparse, "os", SimpleNamespace(name="nt", environ=os.environ))
    def taskkill(command, **kwargs):
        assert command[1:] == ["/PID", "12345", "/T", "/F"]
        if failure == "missing":
            raise FileNotFoundError("mock taskkill unavailable")
        if failure == "timeout":
            raise subprocess.TimeoutExpired("taskkill", 10)
        return SimpleNamespace(returncode=1)
    monkeypatch.setattr(docparse.subprocess, "run", taskkill)
    docparse._terminate_worker(process)
    assert killed == [True]
